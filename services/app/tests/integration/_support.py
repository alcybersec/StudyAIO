"""Shared builders for integration tests: real documents, a fake AI agent, fake embeddings.

Nothing here talks to a provider. `FakeAgent` is a complete `AgentAdapter`, so a
stage that reaches it exercises the real adapter contract (usage metering, result
dataclasses) without a network or a CLI.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

from app.agents.base import (
    AgentAdapter,
    AnswerResult,
    ClassificationResult,
    ConceptData,
    ConceptExtractionResult,
    ConceptRelationData,
    CourseOpsAssessment,
    CourseOpsDeadline,
    CourseOpsResult,
    ExtractionData,
    FlashcardData,
    QuizQuestionData,
    SummaryResult,
)
from app.agents.embeddings import EmbeddingProvider

LECTURE_PAGES = [
    "Networking Fundamentals: the transport layer provides end to end delivery.",
    "TCP uses a three way handshake and congestion control to remain reliable.",
    "UDP is connectionless; applications such as DNS accept loss for speed.",
]


def make_pdf(pages: list[str] | None = None) -> bytes:
    """A real, parseable PDF with one text block per page."""
    import fitz

    doc = fitz.open()
    for text in pages or LECTURE_PAGES:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx(paragraphs: list[str] | None = None) -> bytes:
    """A real DOCX with a heading and body paragraphs."""
    import docx

    document = docx.Document()
    document.add_heading("Lecture 3: Transport", level=1)
    for para in paragraphs or LECTURE_PAGES:
        document.add_paragraph(para)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_pptx(slides: list[str] | None = None) -> bytes:
    """A real PPTX with a title and a body text box per slide."""
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[1]
    for text in slides or LECTURE_PAGES:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = text[:30]
        slide.placeholders[1].text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class FixedEmbeddings(EmbeddingProvider):
    """Deterministic 384-d vectors derived from the text, no model needed."""

    @property
    def dimensions(self) -> int:
        return 384

    def embed_texts(self, texts: list[str]) -> list[list[float]]:
        out = []
        for text in texts:
            seed = (sum(map(ord, text)) % 97) / 97.0
            vec = [0.0] * 384
            vec[0] = 1.0
            vec[1] = seed
            out.append(vec)
        return out

    def preload(self) -> None:
        return None


@dataclass
class FakeAgent(AgentAdapter):
    """A complete adapter that records what it was asked and returns canned data."""

    course_code: str = "NET101"
    week: int = 3
    confidence: float = 0.95
    calls: list[tuple[str, dict]] = field(default_factory=list)

    async def classify_lecture(self, text_preview, filename, known_courses):
        self.calls.append(
            (
                "classify",
                {"preview": text_preview, "filename": filename, "known_courses": known_courses},
            )
        )
        self.usage.add(10, 5)
        return ClassificationResult(
            course_code=self.course_code,
            week=self.week,
            title="Transport Layer",
            confidence=self.confidence,
            reasoning="stub",
        )

    async def generate_summary(self, extraction: ExtractionData, existing_summary):
        self.calls.append(
            ("summary", {"pages": len(extraction.pages), "existing": existing_summary})
        )
        self.usage.add(20, 40)
        version = "update" if existing_summary else "new"
        return SummaryResult(
            content_md=f"# NET101 - Week 3: Transport\n\n## Overview\n\nA {version} summary.\n"
        )

    async def generate_flashcards(self, summary, extraction, count):
        self.calls.append(("flashcards", {"count": count}))
        self.usage.add(5, 5)
        return [
            FlashcardData(front="What does TCP provide?", back="Reliable delivery", tags=["tcp"]),
            FlashcardData(front="Is UDP connection-oriented?", back="No", tags=["udp"]),
        ]

    async def generate_quiz(self, summary, extraction, count):
        self.calls.append(("quiz", {"count": count}))
        self.usage.add(5, 5)
        return [
            QuizQuestionData(
                question_type="multiple_choice",
                question="Which protocol does a three way handshake?",
                options=["TCP", "UDP", "ICMP", "ARP"],
                correct_answer="TCP",
                explanation="TCP opens connections with SYN, SYN-ACK, ACK.",
            ),
            QuizQuestionData(
                question_type="short_answer",
                question="Name one protocol that tolerates loss.",
                correct_answer="UDP",
            ),
        ]

    async def answer_question(self, question, context_chunks):
        self.calls.append(("answer", {"question": question, "chunks": len(context_chunks)}))
        self.usage.add(8, 8)
        return AnswerResult(
            answer=f"Stub answer to: {question}",
            citations=[{"chunk": i} for i, _ in enumerate(context_chunks)],
        )

    async def extract_course_ops(self, document_text, course_code, document_type):
        self.calls.append(("course_ops", {"type": document_type, "course": course_code}))
        self.usage.add(8, 8)
        return CourseOpsResult(
            assessments=[
                CourseOpsAssessment(
                    title="Midterm", assessment_type="exam", weight_pct=30.0, weeks_relevant=[1, 2]
                )
            ],
            deadlines=[
                CourseOpsDeadline(
                    title="Assignment 1", due_date="2030-05-01", deadline_type="assignment"
                )
            ],
            course_info={"name": "Networks"},
            confidence=0.9,
        )

    async def extract_concepts(self, text, existing_concepts=None):
        self.calls.append(("concepts", {"existing": existing_concepts}))
        self.usage.add(8, 8)
        return ConceptExtractionResult(
            concepts=[
                ConceptData(name="TCP", description="Reliable transport", category="protocol"),
                ConceptData(name="UDP", description="Datagram transport", category="protocol"),
            ],
            relations=[ConceptRelationData(source="TCP", target="UDP", relation_type="contrasts")],
        )


# ── Seeded study content ──────────────────────────────────────────────


async def seed_course(
    session,
    user_id: str,
    *,
    code: str = "NET101",
    week: int = 3,
    flashcards: int = 2,
    quizzes: int = 2,
    with_chunks: bool = True,
):
    """A course with one processed artifact, a summary, chunks, flashcards and quizzes.

    Rows are flushed (not committed) so they roll back with the test's SAVEPOINT.
    Returns a namespace with the created ORM objects.
    """
    from types import SimpleNamespace

    from app.core.utils import generate_id
    from app.models.artifact import LectureArtifact
    from app.models.chunk import Chunk
    from app.models.course import Course
    from app.models.extraction import Extraction
    from app.models.flashcard import Flashcard
    from app.models.quiz import QuizQuestion
    from app.models.summary import Summary

    course = Course(id=generate_id(), code=code, name=f"{code} course", user_id=user_id)
    session.add(course)
    await session.flush()

    artifact = LectureArtifact(
        id=generate_id(),
        user_id=user_id,
        course_id=course.id,
        week=week,
        title=f"{code} week {week}",
        original_filename=f"{code}-w{week}.pdf",
        file_path=f"uploads/{generate_id()}_{code}-w{week}.pdf",
        file_type="pdf",
        sha256=generate_id().replace("-", "") * 2,
        file_size_bytes=100,
        status="processed",
    )
    session.add(artifact)
    await session.flush()

    session.add(
        Extraction(
            id=generate_id(),
            artifact_id=artifact.id,
            manifest_json={
                "pages": [
                    {"page_number": i + 1, "text": text, "slide_title": f"Slide {i + 1}"}
                    for i, text in enumerate(LECTURE_PAGES)
                ]
            },
            extraction_path=f"extractions/{artifact.id}",
            page_count=len(LECTURE_PAGES),
            image_count=0,
        )
    )

    summary = Summary(
        id=generate_id(),
        course_id=course.id,
        week=week,
        content_md=f"# {code} - Week {week}\n\n## Overview\n\nTransport protocols and TCP.\n",
        file_path=f"summaries/{course.id}/Week{week}.md",
        version=1,
        source_artifacts=[artifact.id],
    )
    session.add(summary)

    chunks = []
    if with_chunks:
        provider = FixedEmbeddings()
        vectors = provider.embed_texts(LECTURE_PAGES)
        for i, (text, vec) in enumerate(zip(LECTURE_PAGES, vectors, strict=True)):
            chunk = Chunk(
                id=generate_id(),
                artifact_id=artifact.id,
                stable_id=f"{artifact.id}:{i}",
                text=text,
                page_ref=i + 1,
                slide_title=f"Slide {i + 1}",
                embedding=vec,
            )
            session.add(chunk)
            chunks.append(chunk)

    cards = []
    for i in range(flashcards):
        card = Flashcard(
            id=generate_id(),
            course_id=course.id,
            week=week,
            front=f"{code} question {i}?",
            back=f"{code} answer {i}",
            tags=["tcp"],
            source_artifact_id=artifact.id,
            source_page_ref=1,
        )
        session.add(card)
        cards.append(card)

    questions = []
    for i in range(quizzes):
        mcq = i % 2 == 0
        question = QuizQuestion(
            id=generate_id(),
            course_id=course.id,
            week=week,
            question_type="multiple_choice" if mcq else "short_answer",
            question=f"{code} quiz {i}?",
            options_json=["A", "B", "C", "D"] if mcq else None,
            correct_answer="A" if mcq else "an answer",
            explanation="because",
            source_artifact_id=artifact.id,
            source_page_ref=1,
        )
        session.add(question)
        questions.append(question)

    await session.flush()
    return SimpleNamespace(
        course=course,
        artifact=artifact,
        summary=summary,
        chunks=chunks,
        flashcards=cards,
        questions=questions,
    )
